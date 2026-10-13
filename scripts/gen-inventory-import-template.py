"""Generate public/templates/Inventory_Import_Template.xlsx for /catalog/inventory-import.

Headers must match HEADER_ALIASES in zaam-api src/services/import/inventoryExcelParser.ts.
The parser only reads the FIRST sheet and treats every non-empty row after the header as a
product, so the template sheet must stay first and must not contain hint/description rows.

Usage: python3 scripts/gen-inventory-import-template.py [extra output path ...]
"""

import shutil
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'public' / 'templates' / 'Inventory_Import_Template.xlsx'

# (header, required, width, text format, description, allowed values / format, example)
COLUMNS = [
    ('SKU / Item Code', True, 22, True,
     'Unique product code. Used to match existing products (Duplicate SKUs: Skip or Update).',
     'Text. Must be unique per product.', 'SAMPLE-001'),
    ('Barcode / UPC', False, 18, True,
     'Primary barcode. Skipped silently if the barcode already exists on any product.',
     'Digits as text (keep leading zeros). 12 digits = UPC, otherwise EAN.', '2000000000015'),
    ('Product Name', True, 46, False,
     'Product name shown in the catalogue.', 'Text.', 'Sample Disposable Vape 600 Puffs - Blue Razz 20mg'),
    ('Description', False, 46, False, 'Long description.', 'Text.', ''),
    ('Category', False, 20, False, 'Catalogue category.', 'Text.', 'Disposable Vapes'),
    ('Sub-Category', False, 20, False, 'Catalogue sub-category.', 'Text.', '600 Puffs'),
    ('Brand / Manufacturer', False, 20, False, 'Saved as both brand and manufacturer.', 'Text.', 'Sample Brand'),
    ('Unit of Measure (UOM)', False, 14, False, 'Selling unit.', 'Text, e.g. PCS, BOX, PACK.', 'PCS'),
    ('Pack Size', False, 12, True, 'Units per pack.', 'Text or number, e.g. 1, 10, 5 x 10ml.', '10'),
    ('Opening Quantity', False, 12, False,
     'Stock on hand. Only used when Warehouse is filled and "Also import warehouse / stock" is ticked.',
     'Whole number (decimals are rounded).', '50'),
    ('Reorder Level', False, 12, False, 'Stock reorder point (stock import only).', 'Whole number.', '10'),
    ('Reorder Quantity', False, 12, False, 'Quantity to reorder (stock import only).', 'Whole number.', '40'),
    ('Unit Cost Price', False, 12, False, 'Cost price in GBP.', 'Number, no currency symbol.', '2.10'),
    ('Unit Selling Price', False, 12, False, 'Retail selling price in GBP.', 'Number, no currency symbol.', '5.99'),
    ('Tax / GST %', False, 10, False,
     'VAT rate. Matched to an existing tax code (VAT20, VAT0...) or a new one is created.',
     '20 or 0.2 both mean 20%. 0 = zero-rated.', '20'),
    ('Warehouse / Location', False, 30, False,
     'Stock location. Matched by warehouse name or address; a new warehouse is created if no match.',
     'Exact warehouse name as shown in Inventory > Warehouses.', 'Unit 1-2 Chepstow House, ST1 5AJ'),
    ('Bin / Rack No.', False, 12, True, 'Bin in the warehouse; created if missing (stock import only).', 'Text.', 'A-01'),
    ('Supplier Name', False, 22, False,
     'Supplier. Created if missing (stock import only); also saved on the product.', 'Text.', 'Sample Supplier Ltd'),
    ('Supplier SKU', False, 16, True, "Supplier's own product code.", 'Text.', 'SUP-SAMPLE-001'),
    ('Lead Time (Days)', False, 10, False, 'Supplier lead time.', 'Whole number of days.', '7'),
    ('Batch / Lot No.', False, 16, True, 'Batch / lot number (saved on the product and stock row).', 'Text.', 'LOT-SAMPLE-01'),
    ('Expiry Date', False, 13, True,
     'Expiry of the stock batch (stock row only).',
     'UK date dd/mm/yyyy, e.g. 31/12/2027 (also dd-mm-yyyy, dd.mm.yyyy, YYYY-MM-DD or a real Excel date). '
     'mm/yyyy means the last day of that month. Impossible dates (31/02/2027) are left blank with a warning.',
     '31/12/2027'),
    ('Status (Active/Inactive)', False, 14, False,
     'Product status. Blank = Active.', 'Active, Inactive or Discontinued.', 'Active'),
    ('Date Added', False, 13, True,
     'Informational. Read by the importer but not currently saved on the product.',
     'UK date dd/mm/yyyy, e.g. 13/10/2026.', '13/10/2026'),
    ('Remarks', False, 34, False, 'Internal notes saved on the product.', 'Text.', ''),
]

SAMPLE_ROWS = [
    ['SAMPLE-001', '2000000000015', 'Sample Disposable Vape 600 Puffs - Blue Razz 20mg',
     'SAMPLE ROW - delete before importing | 2ml disposable | 20mg nic salt | Blue Razz',
     'Disposable Vapes', '600 Puffs', 'Sample Brand', 'PCS', '10', 50, 10, 40, 2.10, 5.99, 20,
     'Unit 1-2 Chepstow House, ST1 5AJ', 'A-01', 'Sample Supplier Ltd', 'SUP-SAMPLE-001', 7,
     'LOT-SAMPLE-01', '31/12/2027', 'Active', '13/10/2026', 'Sample row - delete before importing'],
    ['SAMPLE-002', '2000000000022', 'Sample E-Liquid 10ml - Strawberry Ice 10mg',
     'SAMPLE ROW - delete before importing | 10ml bottle | 50/50 VG/PG',
     'E-Liquids', 'Nic Salts', 'Sample Brand', 'PCS', '1', 120, 24, 100, 0.85, 3.99, 20,
     'Unit 1-2 Chepstow House, ST1 5AJ', 'B-03', 'Sample Supplier Ltd', 'SUP-SAMPLE-002', 5,
     'LOT-SAMPLE-02', '30/06/2028', 'Active', '13/10/2026', 'Sample row - delete before importing'],
    ['SAMPLE-003', '', 'Sample Mesh Coils 0.8ohm (Pack of 5)',
     'SAMPLE ROW - delete before importing | No barcode, no stock: product only',
     'Accessories', 'Coils', 'Sample Brand', 'PACK', '5', None, None, None, 4.50, 9.99, 20,
     '', '', '', '', None, '', '', 'Active', '13/10/2026', 'Sample row - only the required columns are mandatory'],
    ['SAMPLE-004', '2000000000039', 'Sample Hemp Gummies 300mg (30 Gummies)',
     'SAMPLE ROW - delete before importing | 30 gummies | 10mg each',
     'Health & Wellness', 'CBD Edibles', 'Sample Wellness', 'PCS', '1', 24, 6, 24, 6.00, 14.99, 0,
     '16-22 Hillchurch Street, ST1 2EX', 'C-12', 'Sample Wellness Ltd', 'SW-GUM-300', 14,
     'LOT-SAMPLE-04', '31/03/2027', 'Active', '13/10/2026', 'Sample row - zero-rated example (Tax 0)'],
    ['SAMPLE-005', '2000000000046', 'Sample Pod Kit - Discontinued Colour',
     'SAMPLE ROW - delete before importing | Inactive example',
     'Vape Kits', 'Pod Kits', 'Sample Brand', 'PCS', '1', 0, 0, 0, 8.25, 19.99, 0.2,
     '', '', '', '', None, '', '', 'Inactive', '13/10/2026', 'Sample row - inactive example (Tax as 0.2 = 20%)'],
]

HEADER_FILL = PatternFill('solid', fgColor='1F4E78')
REQUIRED_FILL = PatternFill('solid', fgColor='C00000')
SAMPLE_FILL = PatternFill('solid', fgColor='FFF2CC')
WHITE_BOLD = Font(bold=True, color='FFFFFF')


def build() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Inventory'

    for col, (header, required, width, text_fmt, desc, allowed, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = WHITE_BOLD
        cell.fill = REQUIRED_FILL if required else HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        note = f"{'REQUIRED' if required else 'Optional'}. {desc}\nFormat: {allowed}"
        cell.comment = Comment(note, 'Zaam ERP', width=320, height=140)
        ws.column_dimensions[cell.column_letter].width = width

    for r, values in enumerate(SAMPLE_ROWS, start=2):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=None if value == '' else value)
            cell.fill = SAMPLE_FILL
            if COLUMNS[col - 1][3]:
                cell.number_format = '@'

    for col, spec in enumerate(COLUMNS, start=1):
        if spec[3]:
            letter = ws.cell(row=1, column=col).column_letter
            for r in range(len(SAMPLE_ROWS) + 2, 1001):
                ws[f'{letter}{r}'].number_format = '@'

    status_col = ws.cell(row=1, column=[c[0] for c in COLUMNS].index('Status (Active/Inactive)') + 1).column_letter
    dv = DataValidation(type='list', formula1='"Active,Inactive,Discontinued"', allow_blank=True)
    dv.add(f'{status_col}2:{status_col}1000')
    ws.add_data_validation(dv)

    ws.freeze_panes = 'B2'
    ws.row_dimensions[1].height = 32

    ins = wb.create_sheet('Instructions')
    lines = [
        ['Zaam ERP - Inventory Import template'],
        [],
        ['How to use'],
        ['1. Fill the "Inventory" sheet (it must stay the FIRST sheet). Row 1 is the header row: do not rename, merge or add rows above it.'],
        ['2. Delete the yellow SAMPLE rows before importing. Every non-empty row is imported as a product.'],
        ['3. Required columns (red headers): SKU / Item Code and Product Name. Rows missing either are skipped with a warning.'],
        ['4. In ERP go to Catalog > Inventory Import, choose the file, click Preview and check the warnings and sample, then click Import into system.'],
        ['5. "Duplicate SKUs": Skip existing (the default) leaves products with the same SKU untouched; Update existing overwrites them.'],
        ['6. Quantity, reorder, bin, batch and expiry are only stored when Warehouse / Location is filled and "Also import warehouse / stock" is ticked.'],
        ['7. Column order does not matter, and extra columns are ignored. Header names are matched case-insensitively.'],
        ['8. Dates are UK day/month/year: dd/mm/yyyy (e.g. 07/10/2026 = 7 October 2026). dd-mm-yyyy, dd.mm.yyyy, 2-digit years (07/10/26), YYYY-MM-DD and real Excel date cells also work. Impossible dates such as 31/02/2026 are left blank and listed as a warning in Preview.'],
        ['9. Duplicates are matched by SKU only. The importer does not merge rows with the same barcode or name under different SKUs.'],
        [],
        ['Column', 'Required', 'What it does', 'Allowed values / format', 'Example'],
    ]
    for line in lines:
        ins.append(line)
    header_row = len(lines)
    for col, (header, required, _, _, desc, allowed, example) in enumerate(COLUMNS):
        ins.append([header, 'Yes' if required else 'No', desc, allowed, example])

    ins['A1'].font = Font(bold=True, size=14)
    ins['A3'].font = Font(bold=True)
    for cell in ins[header_row]:
        cell.font = WHITE_BOLD
        cell.fill = HEADER_FILL
    for r in range(header_row + 1, header_row + 1 + len(COLUMNS)):
        if ins.cell(row=r, column=2).value == 'Yes':
            ins.cell(row=r, column=1).font = Font(bold=True, color='C00000')
            ins.cell(row=r, column=2).font = Font(bold=True, color='C00000')
        for c in range(1, 6):
            ins.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical='top')
    for letter, width in zip('ABCDE', (26, 10, 60, 50, 36)):
        ins.column_dimensions[letter].width = width
    ins.freeze_panes = f'A{header_row + 1}'
    return wb


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    build().save(OUT)
    print(f'Wrote {OUT}')
    for extra in sys.argv[1:]:
        shutil.copyfile(OUT, extra)
        print(f'Copied to {extra}')


if __name__ == '__main__':
    main()
